"""Reproducible mixed-document fixtures; run with bundled artifact Python.

Large DOCX byte sizes come from genuine, visible lossless architecture diagrams,
stored without ZIP recompression. Text length is measured independently.
No trailing bytes, opaque padding, hidden text or answer-only metadata are added.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import io
import json
import math
from pathlib import Path
from zipfile import ZIP_STORED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

PROJECT = "TellerX-Stress50-20260907"
NAMES = [
    "澄海支付",
    "青禾清算",
    "星洲退款",
    "云杉账务",
    "霁月对账",
    "锦川授信",
    "海棠票据",
    "寒松归档",
    "碧湖计费",
    "玄鹭通知",
]
ENGLISH = [
    "Clearsea Pay",
    "Greenfield Clearing",
    "Starport Refund",
    "Cedar Ledger",
    "Moon Reconcile",
    "River Credit",
    "Begonia Bills",
    "Pine Archive",
    "Jade Billing",
    "Heron Notify",
]
ROLES = {"architecture": "架构设计", "interface": "接口规范", "operations": "运维手册"}


def dump(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def section(title, level, text="", table=None):
    return {"title": title, "level": level, "text": text, "table": table}


def group_spec(i, role):
    name, en = NAMES[i - 1], ENGLISH[i - 1]
    number = 8800 + i
    doc_id = f"S50-{number}"
    filename = f"{name}二期_{ROLES[role]}_v2.docx"
    opening = f"{name} / {en} 的当前已批准规范。适用范围为生产环境。业务组编号 {doc_id}。"
    basic = [
        section("适用范围 Scope", 1, opening),
        section(
            "功能模块 Functions",
            1,
            "业务接入层检查租户边界并保留请求标识。Duplicate requests reuse the existing operation result.",
        ),
        section(
            "架构模块 Architecture",
            1,
            f"{name}关联控制策略为 CTL-{number}。接口详情见 {name}二期接口规范，运行告警见 {name}二期运维手册。",
        ),
        section(
            "服务结构 Service Tree",
            2,
            "平台 Platform\n  接入 Gateway\n    签名校验 SignatureValidator\n  编排 Orchestrator\n    异常补偿 CompensationWorker\n  存储 Storage\n    审计归档 AuditStore",
        ),
    ]
    if role == "architecture":
        basic += [
            section("接口模块 Interfaces", 1),
            section("安全策略 Security", 2),
            section("签名验证 Signature", 3),
            section(
                "密钥轮换 Key Rotation",
                4,
                f"{name}生产环境签名算法为 HMAC-SHA256。签名密钥轮换周期为 {31 + i} 天。Secret keys are rotated by the security operator.",
            ),
            section(
                "异常处理 Failure Handling",
                2,
                f"CTL-{number} 的同步处理超时后进入 DLQ-{number} 队列。Audit evidence retention is {180 + i * 7} days.\n"
                f"{name}的审计保留期为 {180 + i * 7} 天。不要用其他业务组参数替代当前策略。",
            ),
        ]
    elif role == "interface":
        basic += [
            section("接口模块 Interfaces", 1),
            section("授权接口 Authorization", 2),
            section(
                "请求协议 Request Contract",
                3,
                f"CTL-{number} 对应 POST /stress50/v2/{i:02d}/authorize。Signature header is X-S50-Signature. 策略拒绝码为 ERR-{number}。",
            ),
            section(
                "重试规则 Retry Policy",
                3,
                f"{name}的生产接口最多重试 {i + 2} 次。每次重试间隔 {i + 1} 秒。不可重试的签名错误直接返回 ERR-{number}。",
            ),
            section(
                "区域参数 Regional Parameters",
                2,
                table=[
                    ["Region 区域", "Timeout 超时毫秒", "Approver 审批角色", "Status 状态"],
                    ["APAC 生产", str(350 + i * 17), f"亚太审批岗{i:02d}", "CURRENT"],
                    ["EU 生产", str(520 + i * 17), f"欧洲审批岗{i:02d}", "CURRENT"],
                    ["APAC 测试", str(100 + i * 17), "测试机器人", "TEST ONLY"],
                ],
            ),
        ]
    else:
        basic += [
            section("运维模块 Operations", 1),
            section("异常处理 Failure Handling", 2),
            section(
                "告警路由 Escalation",
                3,
                f"{name}的 DLQ-{number} 队列积压达到 {50 + i * 3} 条时，通知值班角色 Duty-{number}。升级到人工响应的时限为 {10 + i} 分钟。",
            ),
            section(
                "恢复流程 Recovery",
                3,
                "暂停消费 Pause consumer。验证幂等记录 Verify idempotency。重放失败消息 Replay failed messages。恢复消费 Resume consumer。",
            ),
            section(
                "灾备参数 Disaster Recovery",
                2,
                table=[
                    ["Parameter 参数", "Value 值", "Scope 范围"],
                    ["RPO", f"{i + 2} minutes", "生产 Production"],
                    ["RTO", f"{20 + i} minutes", "生产 Production"],
                ],
            ),
        ]
    return {
        "filename": filename,
        "title": f"{name}二期 {ROLES[role]}",
        "group": i,
        "role": role,
        "language": "mixed",
        "sections": basic,
        "lifecycle": "approved",
    }


def independent_spec(i):
    name = [
        "山岚仓储",
        "潮汐采购",
        "萤火报销",
        "磐石设备",
        "竹影排班",
        "白鹭培训",
        "银杏访客",
        "流光发布",
        "长风容量",
        "松涛审计",
    ][i - 1]
    sections = [
        section("业务说明 Overview", 1, f"{name}是独立运行的内部服务。该规范不引用其他业务组。"),
        section("操作规范 Procedure", 1),
        section(
            "批准规则 Approval",
            2,
            f"{name}每批最多处理 {200 + i * 13} 条记录，审批负责人为 Owner-{9000 + i}。The batch has to pass checksum validation before approval.",
        ),
        section(
            "参数表 Parameters",
            2,
            table=[
                ["Field 字段", "Value 值"],
                ["容量 Capacity", str(200 + i * 13)],
                ["负责人 Owner", f"Owner-{9000 + i}"],
            ],
        ),
    ]
    if i == 8:
        sections += [
            section(
                "发布控制 Release Control", 2, "流光发布生产环境同一版本的发布冻结窗口为 45 分钟。"
            ),
            section(
                "审计核准 Audit Approval", 2, "流光发布生产环境同一版本的发布冻结窗口为 75 分钟。"
            ),
        ]
    if i == 9:
        sections += [
            section(
                "环境配额 Environment Limits",
                2,
                "长风容量生产环境的并发上限为 600，测试环境的并发上限为 60。",
            )
        ]
    if i == 10:
        sections += [
            section("审计要求 Audit", 2, "松涛审计的证据保留期为 730 天。"),
            section("运维确认 Operations", 2, "松涛审计要求保存审计证据 730 天。"),
        ]
    return {
        "filename": f"{name}_操作说明.docx",
        "title": f"{name} 操作说明",
        "group": None,
        "role": "independent",
        "language": "mixed",
        "sections": sections,
        "lifecycle": "approved",
    }


def style(doc):
    s = doc.sections[0]
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.top_margin = s.bottom_margin = Inches(0.7)
    s.left_margin = s.right_margin = Inches(0.8)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.color.rgb = RGBColor(0, 0, 0)
        fonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if attr.endswith(("Theme", "theme")):
                del fonts.attrib[attr]
        fonts.set(qn("w:eastAsia"), "PingFang SC")
        for border in st.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        st.font.size = Pt(11 if name == "Normal" else 23 if name == "Title" else 15)
        st.paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.1


def add_table(doc, rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "D9EAF7")
                cell._tc.get_or_add_tcPr().append(shade)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
        if ri == 0:
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    doc.add_paragraph()


def diagram(target_bytes, title):
    width = max(420, int(math.sqrt(target_bytes / 2)))
    height = int(width * 2 / 3)
    image = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(image)
    font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", max(14, width // 40))
    nodes = [
        (0.34, 0.04, 0.66, 0.17, "Platform"),
        (0.06, 0.35, 0.32, 0.5, "Gateway"),
        (0.36, 0.35, 0.65, 0.5, "Orchestrator"),
        (0.71, 0.35, 0.97, 0.5, "Storage"),
        (0.04, 0.74, 0.35, 0.91, "SignatureValidator"),
        (0.37, 0.74, 0.67, 0.91, "CompensationWorker"),
        (0.72, 0.74, 0.97, 0.91, "AuditStore"),
    ]
    for cx, cy in [(0.19, 0.35), (0.50, 0.35), (0.84, 0.35)]:
        d.line(
            [(width * 0.5, height * 0.17), (width * cx, height * cy)],
            fill="#506680",
            width=max(1, width // 450),
        )
    for cx in [0.19, 0.50, 0.84]:
        d.line(
            [(width * cx, height * 0.5), (width * cx, height * 0.74)],
            fill="#506680",
            width=max(1, width // 450),
        )
    for x1, y1, x2, y2, label in nodes:
        d.rectangle(
            (int(width * x1), int(height * y1), int(width * x2), int(height * y2)),
            fill="#E6EEF6",
            outline="#263F59",
            width=max(1, width // 450),
        )
        # Short English component labels avoid font fallback inside the raster.
        box = d.textbbox((0, 0), label, font=font)
        d.text(
            (
                (x1 + x2) * width / 2 - (box[2] - box[0]) / 2,
                (y1 + y2) * height / 2 - (box[3] - box[1]) / 2,
            ),
            label,
            fill="black",
            font=font,
        )
    data = io.BytesIO()
    image.save(data, format="PNG", compress_level=0)
    return data.getvalue()


def write_docx(spec, index, out):
    doc = Document()
    style(doc)
    doc.add_paragraph(spec["title"], "Title")
    doc.add_paragraph(
        "本文件供开发、支持和值班人员查询服务约束。正文说明当前规则、适用环境以及异常情况下的操作顺序。"
    )
    for item in spec["sections"]:
        doc.add_heading(item["title"], item["level"])
        if item["text"]:
            doc.add_paragraph(item["text"])
        if item["table"]:
            add_table(doc, item["table"])
    # Distinct sizes of searchable content as well as physical attachment size.
    count = [0, 4, 12, 24][index % 4]
    for n in range(count):
        doc.add_page_break()
        doc.add_heading(f"运行参考 Runbook {n + 1:02d}", 1)
        for k, (topic, en) in enumerate(
            [
                ("接入核验", "Ingress validation"),
                ("事件追踪", "Event tracing"),
                ("容量计划", "Capacity planning"),
                ("回滚检查", "Rollback review"),
            ],
            1,
        ):
            doc.add_heading(f"{topic} {en}", 2)
            doc.add_paragraph(
                f"针对 {spec['title']} 的第 {n + 1} 组运行案例，{topic}由当班人员执行。"
                f"记录编号 RUN-{index + 1:02d}-{n + 1:03d}-{k}，该编号只用于运行案例检索。"
                "先读取当前配置快照，再比对请求状态和审计流水。出现记录不一致时保留原始证据，"
                "由对应服务团队处理；禁止将另一环境的参数复制到当前运行实例。"
                f"{en} verifies the active configuration and request ledger. Preserve the original event sequence, "
                "record the observed state and ask the owning service to investigate mismatches before resuming work."
            )
    # Needle at the end of the longest files tests position, not filename alone.
    doc.add_heading("末尾恢复检查 Final Recovery", 1)
    token = f"RECOVERY-{9600 + index}"
    doc.add_paragraph(
        f"{spec['title']}的末尾恢复校验口令为 {token}。This check applies after the final replay batch."
    )
    doc.add_page_break()
    doc.add_heading("组件结构图 Component Structure", 1)
    doc.add_paragraph(
        "图中的组件关系与可检索文字结构树一致。CompensationWorker belongs to Orchestrator; AuditStore belongs to Storage."
    )
    target = [512_000, 1_000_000, 2_000_000, 3_000_000, 4_000_000, 5_000_000][index % 6]
    # Keep image members uncompressed so output byte size is reproducible and
    # remains tied to a real image. XML retains normal DOCX compression.
    png = diagram(target - 25_000, spec["title"])
    doc.add_picture(io.BytesIO(png), width=Inches(6.7))
    raw = io.BytesIO()
    doc.save(raw)
    destination = out / spec["filename"]
    with ZipFile(io.BytesIO(raw.getvalue())) as src, ZipFile(destination, "w") as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if info.filename.startswith("word/media/"):
                info.compress_type = ZIP_STORED
            dst.writestr(info, payload)
    spec.update(
        bytes=destination.stat().st_size,
        appendix_pages=count,
        recovery_token=token,
        text_characters=sum(len(p.text) for p in doc.paragraphs),
        target_bytes=target,
    )
    spec["sha256"] = hashlib.sha256(destination.read_bytes()).hexdigest()


def build_questions(specs):
    questions = []

    def q(kind, question, files, values=(), status="answered", heading=None, scope=None):
        questions.append(
            {
                "id": f"q{len(questions) + 1:03d}",
                "kind": kind,
                "question": question,
                "expected_files": files,
                "values": list(values),
                "status": status,
                "heading": heading,
                "scope": scope,
            }
        )

    for i in range(1, 11):
        name = NAMES[i - 1]
        en = ENGLISH[i - 1]
        n = 8800 + i
        a, f, o = specs[(i - 1) * 3 : i * 3]
        q(
            "filename_heading",
            f"{name}二期架构设计文档中，接口模块的安全策略下密钥轮换周期是多少天？",
            [a["filename"]],
            [str(31 + i)],
            heading="密钥轮换",
            scope=a["filename"],
        )
        q(
            "filename_content",
            f"帮我看看{name}二期接口规范这份文档，生产接口最多重试几次？",
            [f["filename"]],
            [str(i + 2)],
            scope=f["filename"],
        )
        q(
            "full_filename",
            f"{o['filename']}文档里，队列积压达到多少条时通知谁？",
            [o["filename"]],
            [str(50 + i * 3), f"Duty-{n}"],
            scope=o["filename"],
        )
        q(
            "table",
            f"{name}二期接口规范文档里，APAC生产的超时和审批角色是什么？",
            [f["filename"]],
            [str(350 + i * 17), f"亚太审批岗{i:02d}"],
            heading="Regional Parameters",
            scope=f["filename"],
        )
        q(
            "global_single",
            f"{name}的签名算法和密钥轮换周期是什么？",
            [a["filename"]],
            ["HMAC-SHA256", str(31 + i)],
        )
        q(
            "cross_document",
            f"{name}同步处理超时后进入哪个队列，积压多少条会通知哪个值班角色？",
            [a["filename"], o["filename"]],
            [f"DLQ-{n}", str(50 + i * 3), f"Duty-{n}"],
        )
        q(
            "global_english",
            f"For {en}, what is the audit evidence retention in days?",
            [a["filename"]],
            [str(180 + i * 7)],
        )
        q(
            "structure_tree",
            f"{name}二期架构设计文档中，CompensationWorker位于哪个父组件下面？",
            [a["filename"]],
            ["Orchestrator"],
            scope=a["filename"],
        )
        q("ambiguous", f"{name}二期文档里有什么规则？", [], status="clarification_required")
        q(
            "scope_insufficient",
            f"{name}二期架构设计文档中，APAC生产接口的超时毫秒数是多少？",
            [],
            status="insufficient_evidence",
            scope=a["filename"],
        )
    for spec in specs[:40]:
        if spec["appendix_pages"] >= 12:
            q(
                "long_document_tail",
                f"{Path(spec['filename']).stem}文档中，末尾恢复校验口令是什么？",
                [spec["filename"]],
                [spec["recovery_token"]],
                heading="末尾恢复检查",
                scope=spec["filename"],
            )
    for i, spec in enumerate(specs[30:40], 1):
        name = spec["title"].split()[0]
        q(
            "independent_global",
            f"{name}每批最多处理多少条记录，谁负责审批？",
            [spec["filename"]],
            [str(200 + i * 13), f"Owner-{9000 + i}"],
        )
    q(
        "true_conflict",
        "流光发布生产环境同一版本的发布冻结窗口是多少分钟？",
        [specs[37]["filename"]],
        ["45", "75"],
        status="conflict",
    )
    q(
        "different_scope",
        "长风容量生产环境与测试环境的并发上限各是多少？",
        [specs[38]["filename"]],
        ["600", "60"],
    )
    q("consistent_facts", "松涛审计的审计证据需要保留多少天？", [specs[39]["filename"]], ["730"])
    q(
        "missing_document",
        "不存在的紫晶项目架构文档里，签名算法是什么？",
        [],
        status="insufficient_evidence",
    )
    q(
        "missing_fact",
        "澄海支付二期接口规范文档里，CEO的私人手机号码是什么？",
        [],
        status="insufficient_evidence",
        scope=specs[1]["filename"],
    )
    for spec in specs[40:]:
        q(
            "simple_document_global",
            f"{spec['title']}的支持时段和联系角色是什么？",
            [spec["filename"]],
            [spec["support"], spec["contact"]],
        )
    return questions


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("evaluation/generated/stress50"))
    args = parser.parse_args()
    root = args.output
    out = root / "documents"
    out.mkdir(parents=True, exist_ok=True)
    specs = [group_spec(i, role) for i in range(1, 11) for role in ROLES]
    specs += [independent_spec(i) for i in range(1, 11)]
    for index, spec in enumerate(specs):
        write_docx(spec, index, out)
        print(
            json.dumps(
                {
                    "document": index + 1,
                    "bytes": spec["bytes"],
                    "appendix_pages": spec["appendix_pages"],
                }
            ),
            flush=True,
        )
    for i in range(1, 11):
        ext = ["md", "html", "txt"][(i - 1) % 3]
        title = f"独立支持站点{i:02d}"
        support = f"{7 + i:02d}:00至{12 + i:02d}:00"
        contact = f"Support-{9700 + i}"
        content = f"{title}提供工作日支持。支持时段为 {support}，联系角色为 {contact}。Support is available on weekdays only.\n附件审阅仅限已完成登记的工单，紧急事件须保留事件编号。"
        filename = f"{title}_SupportGuide.{ext}"
        payload = (
            f"# {title}\n\n## 服务时间 Service Hours\n\n{content}\n"
            if ext == "md"
            else f"<!doctype html><meta charset='utf-8'><h1>{title}</h1><h2>服务时间 Service Hours</h2><p>{html.escape(content)}</p>"
            if ext == "html"
            else title + "\n" + content
        )
        (out / filename).write_text(payload, encoding="utf-8")
        specs.append(
            {
                "filename": filename,
                "title": title,
                "group": None,
                "role": "support",
                "language": "mixed",
                "lifecycle": "approved",
                "bytes": (out / filename).stat().st_size,
                "support": support,
                "contact": contact,
                "sha256": hashlib.sha256((out / filename).read_bytes()).hexdigest(),
            }
        )
    qs = build_questions(specs)
    dump(
        root / "manifest.json",
        {
            "project": PROJECT,
            "documents": specs,
            "count": len(specs),
            "size_method": "Visible lossless PNG diagram stored without ZIP recompression; text lengths vary separately.",
        },
    )
    dump(root / "questions.json", qs)
    print(
        json.dumps(
            {
                "files": len(specs),
                "docx": 40,
                "questions": len(qs),
                "total_bytes": sum(s["bytes"] for s in specs),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
